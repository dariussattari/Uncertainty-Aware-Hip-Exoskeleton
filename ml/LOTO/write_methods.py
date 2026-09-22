"""Write the leave-one-task-out sections into Methods_Experiments_Reproduction.docx.

    python ml/LOTO/write_methods.py --dry-run
    python ml/LOTO/write_methods.py

Idempotent: the sections are delimited by the first new heading and the `References` heading
that follows them, and any previous copy is removed before a fresh one is inserted. Re-run it
whenever more cells finish and the results tables will update in place.

Methodology is written here; the weekly update gets the findings. The two overlap only where a
method detail is needed to read a number honestly.

Style is matched to the surrounding document: `Heading 1` as "N  Title", `Heading 2` as
"N.M  Title", body text in `Normal`, and each table followed by a `Normal` paragraph beginning
"Table NN." Tables in the existing document run to 44, so these start at 45.
"""

from __future__ import annotations

import argparse
import json
import subprocess
import sys
from pathlib import Path

import pandas as pd

HERE = Path(__file__).resolve().parent
REPO = next(d for d in HERE.parents if (d / "data").is_dir())
DOC = Path("/Users/dariussattari/Library/CloudStorage/OneDrive-HarvardUniversity/thesis/"
           "Methods_Experiments_Reproduction.docx")

FIRST_HEADING = "34  Experiment 8: leave-one-task-out"
STOP_HEADING = "References"
TABLE_START = 45
FIG_WIDTH_IN = 6.3

MODE_NAME = {"LG": "level ground", "RA": "ramp ascent", "RD": "ramp descent",
             "SA": "stair ascent", "SD": "stair descent", "ST": "standing",
             "TR": "transitions"}
LEAF_NAME = {"vanilla/AE": "Autoencoder (Exp. 2)", "vanilla/GAN": "GAN (Exp. 3)",
             "vanilla/Ensemble_GP": "Gait-phase ensemble (Exp. 1)",
             "vanilla/Ensemble_synthetic": "Correlation ensemble (Exp. 4)",
             "vanilla/Ensemble_forecast": "Forecast ensemble (Exp. 5)",
             "prob_ensemble": "Probabilistic ensemble (Exp. 7)"}


def gather() -> dict:
    subprocess.run([sys.executable, str(HERE / "summarize.py"), "--quiet"],
                   cwd=REPO, check=True, capture_output=True)
    s = pd.read_csv(HERE / "summary_test.csv")
    pm = pd.read_csv(HERE / "per_mode_test.csv")
    prim = s[s["primary"] & s["auroc"].notna()]
    loto = prim[prim["held_out"] != "participant-split"]
    base = prim[prim["held_out"] == "participant-split"].set_index("leaf")
    cfg = {}
    for fam in ("AE_GAN", "ensemble"):
        f = REPO / "data" / "processed_all" / fam / "config.json"
        if f.exists():
            cfg[fam] = json.loads(f.read_text())
    return {"loto": loto, "base": base, "per_mode": pm, "config": cfg,
            "tasks": [t for t in ("SA", "SD", "ST", "TR") if t in set(loto["held_out"])],
            "models": [m for m in LEAF_NAME if m in set(loto["leaf"])]}


# ---------------------------------------------------------------- the content

def build(g: dict) -> list[tuple]:
    """(kind, payload) blocks in document order."""
    tasks, models = g["tasks"], g["models"]
    n = TABLE_START
    blocks: list[tuple] = []
    P = lambda t: blocks.append(("p", t))
    H1 = lambda t: blocks.append(("h1", t))
    H2 = lambda t: blocks.append(("h2", t))

    def TBL(head, rows, cap):
        nonlocal n
        blocks.append(("tbl", (head, rows)))
        blocks.append(("cap", f"Table {n}. {cap}"))
        n += 1

    # ---------------------------------------------------------------- 34
    H1(FIRST_HEADING)
    P("Every experiment to this point holds out participants. Training used level ground, "
      "ramp ascent and ramp descent from twelve participants; the validation and test "
      "participants contributed those same three modes as in-distribution and four unseen "
      "ones — stair ascent, stair descent, standing and transitions — as pseudo-OOD. This "
      "experiment changes the task partition and nothing else.")

    H2("34.1  What the participant split cannot separate")
    P("The participant-split protocol withholds four tasks simultaneously and trains on none "
      "of them. Under that design no task ever appears as in-distribution in one run and "
      "out-of-distribution in another, so two explanations for a high uncertainty score are "
      "indistinguishable: the task is unfamiliar to the model, or the task simply produces an "
      "unusual signal. Both predict exactly the same measurement.")
    P("Separating them requires the same task to appear on both sides of the comparison. "
      "Holding out one task at a time and training on the remainder does that: each task is "
      "scored once as the withheld class, and again in the runs where it was part of "
      "training. A score that measures unfamiliarity must be higher in the first case. A "
      "score that measures an intrinsic property of the signal will be unchanged.")
    P("The design is also strictly harder than its predecessor, because the held-out task now "
      "has a near neighbour in training. A model asked to flag stair ascent has already "
      "learned stair descent.")

    H2("34.2  The task partition")
    TBL(["Held-out task", "Trained on", "What the run tests"],
        [["Stair ascent (SA)", "LG RA RD SD ST TR",
          "The hard case: stair descent, its nearest neighbour, is in training"],
         ["Standing (ST)", "LG RA RD SA SD TR",
          "The easy case, and the smallest — see the limitation in Section 36"],
         ["Transitions (TR)", "LG RA RD SA SD ST",
          "Mode transitions, with every steady-state mode learned"]],
        "The three task partitions. The participant split is unchanged from Section 4 — the "
        "same seed-0 assignment of twelve training, four validation and four test "
        "participants — so a run here and its participant-split counterpart differ in exactly "
        "one respect. Stair descent is deliberately absent as a holdout; it is the natural "
        "fourth run and is noted as future work.")

    H2("34.3  The windowed dataset had to be rebuilt")
    P("The stored arrays could not support this experiment. The in-distribution definition had "
      "been applied at build time rather than at load time, so the twelve training "
      "participants had no stair, standing or transition windows on disk at all. Holding out "
      "stair ascent requires training on stair descent, standing and transitions; those "
      "windows had never been extracted.")
    P("The arrays were therefore re-extracted from the raw trials, mode-complete, with no "
      "in-distribution or out-of-distribution label stored on disk. The task partition is now "
      "applied when the data is loaded, which is what allows one build to serve all three "
      "holdouts. Window definition, stride, channel order, the participant assignment and the "
      "input scaler are all unchanged; the scaler is copied from the original build rather "
      "than refitted, so scaling is numerically identical to every earlier experiment.")
    TBL(["Set", "Split", "Windows", "Reproduces the original array"],
        [[fam, split, f"{cnt:,}", "yes — window for window"]
         for fam, split, cnt in (("Unlabelled (AE, GAN, label-free ensembles)", "train", 249316),
                                 ("Unlabelled", "validation", 87694),
                                 ("Unlabelled", "test", 84999),
                                 ("Labelled (gait-phase ensemble)", "train", 204561),
                                 ("Labelled", "validation", 75102),
                                 ("Labelled", "test", 72428))],
        "The rebuilt arrays. Restricted to the original three in-distribution modes, each new "
        "split reproduces the corresponding original array exactly — 187,993 / 64,350 / 61,956 "
        "unlabelled and 168,956 / 59,952 / 58,045 labelled windows, verified by comparing "
        "(subject, trial, window-end) keys. The rebuild was repeated independently on the "
        "compute cluster and confirmed bit-identical by hashing the same sampled rows on both "
        "machines.")

    H2("34.4  Implementation")
    P("No model, trainer or evaluator was modified. The task partition is applied by wrapping "
      "each data split in a view that maps row i of the view to row keep[i] of the shared "
      "memory-mapped array; the training split keeps the six trained modes, and the held-out "
      "mode of the validation and test splits becomes the out-of-distribution set. The view "
      "presents the ordinary interface, so the two-stage protocol, the evaluator and the "
      "figures operate unchanged and cannot tell that a holdout occurred. Nothing is copied.")
    P("This matters for attribution. Because the architecture, optimiser, filter, threshold "
      "rule and evaluation code are literally the same objects used in Sections 17 to 33, any "
      "difference between a result here and its participant-split counterpart is attributable "
      "to the task partition alone.")

    H2("34.5  Execution")
    P("Eighteen cells — six models by three holdouts — were run on the FASRC Cannon cluster as "
      "four job arrays: the two single-stage models, the three vanilla ensembles (each running "
      "its own twelve leave-one-subject-out folds in process), the probabilistic ensemble's "
      "thirty-six folds as independent tasks, and its three stage-2 jobs dependent on those. "
      "Sequentially the matrix is roughly 108 hours; as arrays it completes in approximately "
      "the time of one cell.")
    P("Two execution details are recorded because they affect reproducibility rather than "
      "convenience. First, training memory-maps the windows and gathers approximately one "
      "thousand random rows per batch, which on the cluster's networked home filesystem ran at "
      "4.8 MB/s and took 221.8 s per autoencoder epoch against 8 s on a local disk; the arrays "
      "are therefore staged to node-local storage at job start, which restores the epoch time "
      "to 8.4 s. Second, TF32 arithmetic was deliberately left disabled. It reduces an epoch "
      "from 16.4 s to 6.6 s on an A100, but truncates the mantissa of every convolution, and "
      "these results are read against participant-split results computed in full precision.")

    # ---------------------------------------------------------------- 35
    H1("35  Deviations — leave-one-task-out")
    TBL(["Item", "Kind", "Detail"],
        [["Task partition", "extension",
          "Not in the reference work. One ambulation mode is withheld and the remaining six "
          "are trained on, in place of withholding four and training on three."],
         ["Out-of-distribution prevalence", "consequence",
          "Falls from 27.1% of test windows to 12.1% (SA), 2.3% (TR) and 0.7% (ST). Precision, "
          "recall at a fixed threshold and F1 all move with prevalence and are therefore not "
          "comparable across protocols; AUROC and best achievable J are."],
         ["Stored ID/OOD labels", "correction",
          "Removed. The partition is applied at load time, which is what permits one build to "
          "serve several holdouts and removes a class of silent error described in Section 36."],
         ["Held-out task set", "scope",
          "Three of the four available modes. Stair descent is not used as a holdout in this "
          "round."]],
        "Deviations introduced by this experiment, in the format of Sections 18, 23, 26 and 31.")

    # ---------------------------------------------------------------- 36
    H1("36  Verification, and a fault it exposed")
    P("Three figure modules constructed their own data object instead of receiving the one the "
      "model was trained and scored with. That object reads the original arrays, whose "
      "in-distribution partition is the participant-split one. Under a task holdout the model, "
      "the windows and every uncertainty value were correct, and every number written to the "
      "result files was correct, because the evaluator did receive the right data object. Only "
      "the figures were affected: tasks that were in training were drawn as held-out, the "
      "saved score bundle recorded 23,043 out-of-distribution windows where the metrics "
      "recorded 10,273, and the ROC curves were computed against the wrong ground truth. "
      "Nothing raised an error and every figure looked plausible.")
    P("The fix makes the dependency explicit: the figure entry points now require the data "
      "builder, and the command-line layer passes the one the model was trained with, so "
      "agreement is structural rather than conventional. Plot titles name the modes actually "
      "withheld in that run, read from the labels rather than from a constant.")
    P("A verification script now checks every result directory against the one fact that "
      "defines it, its own name: that the evaluation marks exactly the named task as "
      "out-of-distribution, that its window counts match the rebuilt arrays, that the saved "
      "score bundle agrees with the metrics window for window, and that no figure predates the "
      "bundle it was drawn from. Reintroducing the original fault into a copied directory "
      "makes the script fail with all three relevant messages, so it is known to detect it.")

    # ---------------------------------------------------------------- 37
    H1("37  Results — leave-one-task-out")
    H2("37.1  Detection of the held-out task")
    head = ["Model", "Participant split"] + [f"Held out {t}" for t in tasks]
    rows = []
    for leaf in models:
        sub = g["loto"][g["loto"]["leaf"] == leaf]
        b = g["base"]["auroc"].get(leaf)
        r = [LEAF_NAME[leaf], f"{b:.3f}" if pd.notna(b) else "—"]
        for t in tasks:
            v = sub[sub["held_out"] == t]["auroc"]
            r.append(f"{v.iloc[0]:.3f}" if len(v) else "not yet run")
        rows.append(r)
    TBL(head, rows,
        "Area under the ROC curve for identifying the held-out task, under the evaluation "
        "procedure of Section 7. The participant-split column is the corresponding result from "
        "Sections 19, 24, 27 and 33 and is shown for orientation only: its test set is 27.1% "
        "out-of-distribution against 0.7-12% here, so only AUROC and best achievable J are "
        "comparable between the columns.")

    H2("37.2  How the scores respond to training membership")
    P("The measurement this experiment exists to make is the one below: the uncertainty "
      "assigned to a task in the run that withheld it, against the same quantity in the runs "
      "that trained on it. A score that measures unfamiliarity must be higher in the first "
      "case, for every task and every model.")
    pm = g["per_mode"]
    rows = []
    for leaf in models:
        sub = pm[pm["leaf"] == leaf]
        for m in ("SA", "ST", "TR"):
            held = sub[(sub["mode"] == m) & (sub["held_out"] == m)]["median_psi"]
            inn = sub[(sub["mode"] == m) & (sub["held_out"] != m)]["median_psi"]
            if not len(held) or not len(inn) or pd.isna(held.iloc[0]) or not float(inn.mean()):
                continue
            h, i = float(held.iloc[0]), float(inn.mean())
            rows.append([LEAF_NAME[leaf], f"{MODE_NAME[m]} ({m})", f"{h:.5g}", f"{i:.5g}",
                         f"{100 * (h - i) / i:+.0f}%"])
    if rows:
        TBL(["Model", "Task", "Median Psi when held out", "Median Psi when trained on",
             "Change"], rows,
            "Median uncertainty for a task, withheld against trained on, each score at its "
            "own scale. Read the sign column by task rather than by model: transitions are "
            "positive in all six models, standing is negative in all six, and stair ascent is "
            "mixed. Only the transitions column has the sign a novelty detector requires.")
    P("The result is not a uniform absence of response. The sign is consistent across models "
      "but is set by the task, which is the opposite of what a working novelty score would "
      "show. Withholding transitions raises the score in every one of the six models, by one "
      "to fifty-three per cent — a genuine if modest novelty response. Withholding standing "
      "LOWERS the score in every one of the six, by three to eighty-one per cent. Stair ascent "
      "is mixed, three models up and three down.")
    P("The standing result has a mechanism, and it is worth stating because it generalises. "
      "Standing is intrinsically ambiguous for every target used here: gait phase is "
      "ill-defined for a stationary limb, and a near-static window carries little to forecast. "
      "Training on standing therefore forces the branches to disagree about it and gives the "
      "autoencoder a latent region of its own, both of which raise the score. Removing "
      "standing from training removes that disagreement. Training on a task can therefore "
      "increase the uncertainty assigned to it, which is precisely backwards for a detector.")
    P("The magnitudes also matter. With the exception of the autoencoder's standing row, these "
      "changes are small beside the spread between tasks within any single run, which is a "
      "factor of two to twenty depending on the model. A run's headline figure is therefore "
      "still dominated by where the withheld task sits in the score's fixed ranking of the "
      "seven tasks, not by whether it was trained on.")

    H2("37.3  The rankings are model-specific and, in one case, inverted")
    P("What separates the estimators is not whether they respond to novelty — none of them "
      "does, materially — but which intrinsic property their ranking encodes. The autoencoder "
      "and the GAN both place standing at the top: it is the most unusual-looking signal in "
      "the set, being close to static. The forecasting ensemble places standing at the bottom "
      "and the stair modes at the top, because a near-static signal is trivial to predict 200 "
      "ms ahead and produces almost no disagreement between branches.")
    P("Those two orderings are close to opposites, which explains the pattern in Section 37.1 "
      "without appealing to detection ability at all. Withholding stair ascent, the "
      "forecasting ensemble is well above chance and the GAN is at chance. Withholding "
      "standing, the GAN is near-perfect and the forecasting ensemble is at chance. Both "
      "results follow from where the withheld task sits in each fixed ranking.")
    P("This is consistent with, and considerably strengthens, the amplitude confound already "
      "recorded in Sections 19.3 and 27.4. Those sections established that a score correlated "
      "with signal amplitude; this experiment establishes that the score does not additionally "
      "respond to familiarity, which the earlier protocol could not test.")

    H2("37.4  One effect that is training-dependent")
    P("Standing's epistemic uncertainty under the probabilistic ensemble was 0.036 in the "
      "participant-split run, where training was level ground and ramps only, and falls to "
      "0.004-0.006 once stairs and transitions are present in training. Additional task "
      "variety therefore does reduce branch disagreement substantially. It reduces it for "
      "withheld standing as much as for trained standing, however, which is the opposite of "
      "the behaviour a novelty score requires.")

    H2("37.5  The inherited threshold does not transfer")
    P("The 99.5th-percentile operating point of Section 7 transfers poorly to this protocol. "
      "Several runs flag no test window at all, and several produce a negative J statistic "
      "because false alarms on in-distribution tasks outnumber true detections of the withheld "
      "one. Two mechanisms contribute: out-of-distribution prevalence is far lower here, and "
      "the training distribution now contains the extreme-uncertainty modes that previously "
      "sat only in the held-out set, which moves the percentile itself. Best achievable J is "
      "reported alongside J at the inherited threshold throughout, and fitting the operating "
      "point on validation data is recorded as future work.")

    H2("37.6  Limitations of this experiment")
    P("Standing is thin on both sides. The study pool contains five standing trials in total, "
      "two of which fall in the training participants and one in the test participants, so the "
      "standing holdout removes almost nothing from training and is scored on 621 windows from "
      "a single trial of a single participant. The standing row should be read as a "
      "consistency check rather than a result. Transitions are thin on the training side for "
      "the same reason, contributing under two per cent of training windows when they are not "
      "the withheld task.")
    P("Stair descent was not used as a holdout. It is the same size as stair ascent and "
      "withholding it while stair ascent remains in training is the mirror image of the "
      "hardest run here; it is the obvious next addition.")
    P("Finally, the benchmark limitation of Section 8 is unchanged. Held-out locomotion is "
      "pseudo-OOD: every task in this dataset is cyclic lower-limb ambulation, and none is the "
      "genuinely non-cyclic activity the reference work used.")

    # figures
    for name, cap in ((f"medianpsi_test.png",
                       "Median uncertainty per task, one line per run, rings marking the task "
                       "that run withheld. The rings lie on the line rather than above it. The "
                       "horizontal ordering differs between panels and is the substantive "
                       "result: the autoencoder and GAN rank standing highest, the forecasting "
                       "ensemble ranks it lowest and the stair modes highest."),
                      (f"auroc_test.png",
                       "Detection of the held-out task by model and holdout, with the "
                       "participant-split result in grey for orientation.")):
        f = HERE / name
        if f.exists():
            blocks.append(("fig", (f, cap)))
    return blocks


# ---------------------------------------------------------------- docx surgery

def render(blocks: list[tuple], dry: bool) -> int:
    if dry:
        fign = 1
        for kind, payload in blocks:
            if kind == "h1":
                print(f"\n{'=' * 78}\n{payload}\n{'=' * 78}")
            elif kind == "h2":
                print(f"\n--- {payload}")
            elif kind == "p":
                print(f"  {payload[:300]}")
            elif kind == "tbl":
                head, rows = payload
                print("\n  " + " | ".join(head))
                for r in rows:
                    print("  " + " | ".join(str(x) for x in r))
            elif kind == "cap":
                print(f"  [{payload[:140]}]")
            else:
                print(f"\n  [figure {fign}: {payload[0].name}]"); fign += 1
        return 0

    import docx
    from docx.shared import Inches
    from docx.text.paragraph import Paragraph

    d = docx.Document(DOC)
    body = d.element.body
    children = list(body.iterchildren())
    start = stop = None
    for i, el in enumerate(children):
        if not el.tag.endswith("}p"):
            continue
        p = Paragraph(el, d)
        t = p.text.strip()
        if start is None and t == FIRST_HEADING:
            start = i
        if t == STOP_HEADING and p.style.name.startswith("Heading"):
            stop = i
            break
    if stop is None:
        raise SystemExit(f"could not find the {STOP_HEADING!r} heading")
    if start is not None:
        for el in children[start:stop]:
            body.remove(el)
        print(f"  removed the previous {FIRST_HEADING!r} block")
        children = list(body.iterchildren())
        stop = next(i for i, el in enumerate(children)
                    if el.tag.endswith("}p") and Paragraph(el, d).text.strip() == STOP_HEADING)
    anchor = Paragraph(children[stop], d)

    def insert(el):
        anchor._p.addprevious(el)

    def par(text, style):
        p = d.add_paragraph(text, style=style)
        insert(p._p)
        return p

    for kind, payload in blocks:
        if kind == "h1":
            par(payload, "Heading 1")
        elif kind == "h2":
            par(payload, "Heading 2")
        elif kind in ("p", "cap"):
            par(payload, "Normal")
        elif kind == "tbl":
            head, rows = payload
            tb = d.add_table(rows=1 + len(rows), cols=len(head))
            # Match the document's own tables rather than imposing a style it does not define:
            # every existing table here is "Normal Table". Setting a style the file lacks
            # raises, and adding one would make these tables look unlike their neighbours.
            for want in ("Normal Table", "Table Grid"):
                try:
                    tb.style = want
                    break
                except KeyError:
                    continue
            for j, v in enumerate(head):
                c = tb.cell(0, j)
                c.text = str(v)
                for r in c.paragraphs[0].runs:
                    r.bold = True
            for i, row in enumerate(rows, start=1):
                for j, v in enumerate(row):
                    tb.cell(i, j).text = str(v)
            insert(tb._tbl)
        else:
            path, cap = payload
            p = d.add_paragraph()
            p.add_run().add_picture(str(path), width=Inches(FIG_WIDTH_IN))
            insert(p._p)
    d.save(DOC)
    print(f"\nwrote sections 34-37 into {DOC.name}")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    g = gather()
    if not len(g["loto"]):
        print("no leave-one-task-out results yet")
        return 0
    print(f"{len(g['loto'])} run(s), {len(g['models'])} model(s), tasks {g['tasks']}")
    return render(build(g), a.dry_run)


if __name__ == "__main__":
    raise SystemExit(main())
